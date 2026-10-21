"""
paper/make_docx.py

Converts the compiled LaTeX paper (paper.tex + paper.aux from tectonic) into
paper.docx with pandoc, then fixes up the Word formatting with python-docx.

Pandoc cannot resolve \\ref or number LaTeX floats itself, so cross-reference
numbers are taken from paper.aux and written into the text and captions
before conversion. Figures use the PNG versions in figures/.
"""

import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
tex = (HERE / "paper.tex").read_text()
aux = (HERE / "paper.aux").read_text()

labels = dict(re.findall(r"\\newlabel\{([^}]+)\}\{\{([^}]*)\}", aux))

# ---------------------------------------------------------------------------
# LaTeX -> pandoc-friendly LaTeX
# ---------------------------------------------------------------------------
body = tex[tex.index("\\begin{document}") + len("\\begin{document}"):tex.index("\\end{document}")]
body = body.replace("\\maketitle", "")


def number_captions(text, env, word):
    """Prefix the float's own caption with 'Figure N.' / 'Table N.' and its subcaptions with (a), (b)."""
    out, pos = [], 0
    pat = re.compile(r"\\begin\{" + env + r"\}(\[[^\]]*\])?(.*?)\\end\{" + env + r"\}", re.S)
    for m in pat.finditer(text):
        block = m.group(2)
        lab = re.search(r"\\label\{([^}]+)\}", block)
        n = labels.get(lab.group(1), "?") if lab else "?"
        # subfigure captions get letters
        letters = iter("abcdefgh")
        block = re.sub(r"(\\begin\{subfigure\}.*?)\\caption\{",
                       lambda s: s.group(1) + f"\\caption{{({next(letters)}) ", block, flags=re.S)
        # the float's own caption is the last \caption outside subfigures
        idx = block.rfind("\\caption{")
        block = block[:idx] + f"\\caption{{\\textbf{{{word} {n}.}} " + block[idx + len("\\caption{"):]
        out.append(text[pos:m.start()] + f"\\begin{{{env}}}" + block + f"\\end{{{env}}}")
        pos = m.end()
    out.append(text[pos:])
    return "".join(out)


body = number_captions(body, "figure", "Figure")
body = number_captions(body, "table", "Table")
body = re.sub(r"\\ref\{([^}]+)\}", lambda m: labels.get(m.group(1), "??"), body)
body = body.replace("\\checkmark", "✓")
# run-in paragraph headings instead of numbered level-4 headings
body = re.sub(r"\\paragraph\{([^}]*)\}", r"\\textbf{\1}", body)
body = re.sub(r"\\setlength\{\\tabcolsep\}\{[^}]*\}", "", body)
body = body.replace("\\small\n", "\n")
# subfigures: pandoc renders them, but keep each one as a plain centred figure
body = re.sub(r"\\begin\{subfigure\}(\[[^\]]*\])?\{[^}]*\}", "", body)
body = body.replace("\\end{subfigure}\\hfill", "").replace("\\end{subfigure}", "")
body = re.sub(r"\\includegraphics\[[^\]]*\]\{([^}]+)\}", r"\\includegraphics{\1}", body)
# references where the LaTeX bibliography was; the appendix stays after them
body = body.replace("\\bibliographystyle{unsrtnat}\n\\bibliography{refs}",
                    "\\section*{References}\nREFS-PLACEHOLDER\n")

title = "Does Attention Help Recurrent Models for Wearable Human Activity Recognition?"
subtitle = "A Subject-Independent Ablation with a Pre-Specified Decision Rule on Three Benchmarks"
preamble = r"""\documentclass{article}
\usepackage{amsmath}
\usepackage{graphicx}
\usepackage{booktabs}
\begin{document}
"""
(HERE / "generated" / "pandoc_input.tex").write_text(preamble + body + "\n\\end{document}\n")

(HERE / "generated" / "refs.lua").write_text(r'''
function Para(el)
  if pandoc.utils.stringify(el) == "REFS-PLACEHOLDER" then
    return pandoc.Div({}, {id = "refs"})
  end
end
''')

meta = HERE / "generated" / "meta.yaml"
meta.write_text(f"""---
title: "{title}"
subtitle: "{subtitle}"
author:
  - "Himansh Munjal — [Department, Institution, City, Country] — [corresponding-author e-mail]"
link-citations: true
---
""")

subprocess.run([
    "pandoc", str(HERE / "generated" / "pandoc_input.tex"), "--metadata-file", str(meta),
    "-f", "latex", "-t", "docx",
    "--resource-path", f"{HERE / 'figures'}:{HERE}",
    "--default-image-extension", "png",
    "--number-sections",
    "--lua-filter", str(HERE / "generated" / "refs.lua"),
    "--citeproc", "--bibliography", str(HERE / "refs.bib"), "--csl", str(HERE / "ieee.csl"),
    "-o", str(HERE / "paper.docx"),
], check=True)

# ---------------------------------------------------------------------------
# Word formatting
# ---------------------------------------------------------------------------
doc = Document(HERE / "paper.docx")

for s in doc.sections:
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.5)

SERIF = "Times New Roman"


def set_font(style, size=None, bold=None, color=None):
    style.font.name = SERIF
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(k), SERIF)
    for k in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(k)) is not None:
            del rfonts.attrib[qn(k)]
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


BLACK = RGBColor(0, 0, 0)
for name, size, bold in [("Normal", 11, None), ("Body Text", 11, None), ("First Paragraph", 11, None),
                         ("Compact", 11, None), ("Title", 17, True), ("Subtitle", 13, False), ("Author", 11, False),
                         ("Abstract", 10, None), ("Heading 1", 14, True), ("Heading 2", 12, True),
                         ("Heading 3", 11, True), ("Heading 4", 11, True), ("Image Caption", 9.5, None),
                         ("Table Caption", 9.5, None), ("Bibliography", 9.5, None), ("Abstract Title", 11, True)]:
    try:
        st = doc.styles[name]
    except KeyError:
        continue
    set_font(st, size, bold, BLACK)
    pf = st.paragraph_format
    if name in ("Body Text", "First Paragraph", "Normal", "Compact"):
        pf.space_after = Pt(5)
        pf.line_spacing = 1.12
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if name.startswith("Heading"):
        pf.space_before = Pt(12 if name == "Heading 1" else 8)
        pf.space_after = Pt(4)
        st.font.italic = False
    if name in ("Image Caption", "Table Caption"):
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.space_after = Pt(10)
        st.font.italic = False
    if name in ("Title", "Subtitle", "Author"):
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if name == "Abstract":
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# centre figures and keep them within the text width
TEXT_W = Cm(16.0)
for p in doc.paragraphs:
    for d in p._p.iter(qn("wp:inline")):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ext = d.find(qn("wp:extent"))
        cx, cy = int(ext.get("cx")), int(ext.get("cy"))
        if cx > TEXT_W:
            ratio = TEXT_W / cx
            ext.set("cx", str(int(cx * ratio)))
            ext.set("cy", str(int(cy * ratio)))
            for e in d.iter(qn("a:ext")):
                e.set("cx", str(int(cx * ratio)))
                e.set("cy", str(int(cy * ratio)))


def border(el, edges, sz=8):
    pr = el.find(qn("w:tblPr")) if el.tag == qn("w:tbl") else el.get_or_add_tcPr()
    tag = "w:tblBorders" if el.tag == qn("w:tbl") else "w:tcBorders"
    b = pr.find(qn(tag))
    if b is None:
        b = OxmlElement(tag)
        pr.append(b)
    for edge in edges:
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), "000000")
        b.append(e)


# booktabs-like tables: rule above, below the header row, and at the bottom
for t in doc.tables:
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if any(True for _ in t._tbl.iter(qn("wp:inline"))):
        continue  # layout table pandoc uses for side-by-side subfigures
    border(t._tbl, ["top", "bottom"], sz=10)
    if t.rows:
        for c in t.rows[0].cells:
            border(c._tc, ["bottom"], sz=6)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT if c is row.cells[0] else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    r.font.name = SERIF

# page numbers in the footer
for s in doc.sections:
    fp = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    for kind, text in [("begin", None), (None, "PAGE"), ("end", None)]:
        if kind:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kind)
            run._r.append(f)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            run._r.append(it)

doc.save(HERE / "paper.docx")
print("wrote", HERE / "paper.docx")
